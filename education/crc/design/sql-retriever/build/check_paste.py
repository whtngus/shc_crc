# python check_paste.py {덱}.pptx — 본 덱에 붙여 넣기 전 점검(pptx-guide 6-13)
# 가이드 점검에 더해 빈 폴더 항목, 둥글기 값 위치, 14pt 미만 글자, 가운뎃점, Pretendard 외 글꼴도 확인함
import re
import sys
import zipfile

from pptx import Presentation

p = Presentation(sys.argv[1])
err = []
if (p.slide_width, p.slide_height) != (14630400, 8229600):
    err.append("슬라이드 크기가 16x9인치가 아님")
for i, s in enumerate(p.slides, 1):
    if s.slide_layout.name != "내용_단쪽":
        err.append(f"{i}쪽 레이아웃 이름: {s.slide_layout.name}")
    if s.shapes.title is None or not s.shapes.title.text.strip():
        err.append(f"{i}쪽 제목 개체 틀 비었음")
    for sh in s.shapes:
        if sh.has_text_frame and " / " in sh.text_frame.text and sh.top > 7_500_000 and len(sh.text_frame.text) < 10:
            err.append(f"{i}쪽 글자 쪽 표시 의심: {sh.text_frame.text}")
xml = "".join(s._element.xml for s in p.slides)
if "schemeClr" in xml:
    err.append("테마 색(schemeClr) 사용")

with zipfile.ZipFile(sys.argv[1]) as z:
    names = z.namelist()
    if any(n.endswith("/") for n in names):
        err.append("zip 안에 빈 폴더 항목 있음")
    for n in names:
        if not re.match(r"ppt/slides/slide\d+\.xml$", n):
            continue
        x = z.read(n).decode("utf8")
        # 둥글기 값(adj)은 roundRect에만 있어야 함
        for m in re.finditer(r'<a:prstGeom prst="(\w+)"><a:avLst>(.*?)</a:avLst>', x):
            if m.group(2) and m.group(1) != "roundRect":
                err.append(f"{n}: {m.group(1)}에 둥글기 값")
        small = sorted({int(v) for v in re.findall(r'<a:(?:rPr|defRPr|endParaRPr)[^>]*\ssz="(\d+)"', x) if int(v) < 1400})
        if small:
            err.append(f"{n}: 14pt 미만 글자 {small}")
        if "·" in x:
            err.append(f"{n}: 가운뎃점 있음")
        for f in sorted(set(re.findall(r'typeface="([^"]+)"', x))):
            if not f.startswith("Pretendard"):
                err.append(f"{n}: 글꼴 {f}")
print("\n".join(err) or "OK")
